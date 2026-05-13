"""Excel parser using openpyxl."""

import openpyxl


def parse_excel(file_path: str) -> list[dict]:
    """Parse an Excel file, returning one item per sheet as {text, page_number, metadata}."""
    results = []
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=True):
            # Filter out empty rows, join cells with tab
            row_text = "\t".join(str(c) if c is not None else "" for c in row)
            if row_text.strip():
                rows.append(row_text)

        if rows:
            results.append({
                "text": "\n".join(rows),
                "page_number": None,
                "metadata": {
                    "sheet_name": sheet_name,
                    "file_type": "excel",
                },
            })
    wb.close()
    return results
